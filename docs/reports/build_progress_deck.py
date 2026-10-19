# -*- coding: utf-8 -*-
"""Generate the MAWE progress report deck.

    python docs/reports/build_progress_deck.py

Numbers in this deck come from the test suite and the milestone scripts; see
WEEK4-PROGRESS-REPORT.pptx next to this file.
"""

from __future__ import annotations

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "WEEK4-PROGRESS-REPORT.pptx")

NAVY = RGBColor(0x00, 0x34, 0x66)
LBLUE = RGBColor(0xBB, 0xE0, 0xE3)
PALE = RGBColor(0xED, 0xF5, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
GREEN = RGBColor(0x1E, 0x6F, 0x3C)
AMBER = RGBColor(0xA8, 0x62, 0x00)
RED = RGBColor(0xB0, 0x30, 0x25)

FONT = "Arial"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def textbox(slide, x, y, w, h, lines, size=18, color=BLACK, bold=False,
            align=PP_ALIGN.LEFT, space_after=8, line_spacing=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(lines):
        text, sz, bd, col = (item if isinstance(item, tuple)
                             else (item, size, bold, color))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        if not text:
            continue
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.name = Pt(sz), bd, FONT
        r.font.color.rgb = col
    return box


def header(slide, title, kicker=None):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(W), Inches(1.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False
    textbox(slide, 0.6, 0.22, W - 3.2, 0.7, [(title, 28, True, WHITE)])
    if kicker:
        textbox(slide, W - 3.6, 0.36, 3.0, 0.5, [(kicker, 13, False, LBLUE)],
                align=PP_ALIGN.RIGHT)


def footer(slide, page):
    textbox(slide, 0.6, H - 0.52, 8.0, 0.3,
            [("MAWE Capstone  ·  Progress Report  ·  Week 4 of 12", 10, False, GREY)])
    textbox(slide, W - 1.4, H - 0.52, 0.8, 0.3,
            [(str(page), 10, False, GREY)], align=PP_ALIGN.RIGHT)


def table(slide, x, y, w, rows, widths, size=13, row_h=0.42,
          status_col=None, status_colors=None):
    n, c = len(rows), len(rows[0])
    gf = slide.shapes.add_table(n, c, Inches(x), Inches(y), Inches(w),
                                Inches(row_h * n))
    tbl = gf.table
    total = sum(widths)
    for j, cw in enumerate(widths):
        tbl.columns[j].width = Emu(int(Inches(w) * cw / total))
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.09)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (WHITE if i % 2 else PALE)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(size)
            r.font.name = FONT
            r.font.bold = (i == 0)
            if i == 0:
                r.font.color.rgb = WHITE
            elif status_col is not None and j == status_col and status_colors:
                r.font.color.rgb = status_colors.get(str(val), BLACK)
                r.font.bold = True
            else:
                r.font.color.rgb = BLACK
    return gf


def card(slide, x, y, w, h, title, body, accent=NAVY, title_size=15,
         body_size=12.5):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = PALE
    box.line.color.rgb = accent
    box.line.width = Pt(1.25)
    box.shadow.inherit = False
    box.text_frame.text = ""
    textbox(slide, x + 0.22, y + 0.16, w - 0.44, 0.4,
            [(title, title_size, True, accent)])
    textbox(slide, x + 0.22, y + 0.62, w - 0.44, h - 0.78,
            [(body, body_size, False, BLACK)], line_spacing=1.12)


def bullet_block(slide, x, y, w, items, size=15, gap=10, marker="▪"):
    lines = []
    for it in items:
        lines.append((f"{marker}  {it}", size, False, BLACK))
    return textbox(slide, x, y, w, 0.4 * len(items), lines, space_after=gap,
                   line_spacing=1.1)


# =========================================================== 1 · title slide
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(W), Inches(H))
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
bg.shadow.inherit = False

textbox(s, 1.0, 1.75, 11.3, 1.0,
        [("Multi-Agent AI Workflow Engine", 46, True, WHITE)])
textbox(s, 1.0, 2.75, 11.3, 0.5,
        [("Simulating an Autonomous Software Development Team", 21, False, LBLUE)])

rule = s.shapes.add_shape(1, Inches(1.0), Inches(3.45), Inches(2.4), Inches(0.05))
rule.fill.solid()
rule.fill.fore_color.rgb = LBLUE
rule.line.fill.background()
rule.shadow.inherit = False

textbox(s, 1.0, 3.75, 11.3, 1.4, [
    ("PROGRESS REPORT  ·  END OF WEEK 4 OF 12", 16, True, WHITE),
    ("", 8, False, WHITE),
    ("Marco Perozo  ·  Gabriel Marte  ·  Lewdy Aquino  ·  Alvin Almonte", 15, False, LBLUE),
    ("Department of Computer Science  ·  Southwest Minnesota State University", 13, False, LBLUE),
], space_after=4)

textbox(s, 1.0, 5.75, 11.3, 0.9, [
    ("Repository", 12, True, LBLUE),
    ("github.com/marcoaperozo/Capstone-MAWE", 18, True, WHITE),
], space_after=3)

textbox(s, 1.0, H - 0.75, 11.3, 0.3,
        [("Phase 1 complete  ·  Phase 2 in progress  ·  114 automated tests passing",
          12, False, LBLUE)])

# ===================================================== 2 · where the project is
s = prs.slides.add_slide(BLANK)
header(s, "Where the project stands", "Week 4 of 12")
textbox(s, 0.6, 1.45, 12.1, 0.5, [
    ("Phase 1 is complete and Phase 2 is two thirds done. The engine runs end "
     "to end, every run is stored, and three milestone checks run automatically "
     "on every push.", 16, False, BLACK)], line_spacing=1.15)

table(s, 0.6, 2.35, 12.1, [
    ["Phase", "Weeks", "Deliverable", "Status"],
    ["1 · Setup & prototyping", "1–2",
     "Repository, CI, database schema, pipeline on a state graph", "COMPLETE"],
    ["2 · Role implementation", "3–5",
     "Product Manager, Developer and Code Reviewer agents", "IN PROGRESS"],
    ["3 · Memory & state", "6–8",
     "Context summarisation, token budget, loop hardening", "PLANNED"],
    ["4 · Testing & evaluation", "9–12",
     "Benchmark suite, single-model baseline, final report", "PLANNED"],
], [2.5, 0.8, 5.6, 1.5], size=13, row_h=0.62, status_col=3,
    status_colors={"COMPLETE": GREEN, "IN PROGRESS": AMBER, "PLANNED": GREY})

card(s, 0.6, 5.35, 3.85, 1.5, "Milestone M1 — met",
     "Three-role pipeline runs end to end on a LangGraph state machine, with "
     "every exchange persisted.", GREEN)
card(s, 4.72, 5.35, 3.85, 1.5, "Milestone M2 — on track",
     "Product Manager and Developer complete and tested. The Code Reviewer is "
     "Week 5.", AMBER)
card(s, 8.85, 5.35, 3.85, 1.5, "Schedule",
     "On plan. No milestone has slipped and no scope has been dropped.", NAVY)
footer(s, 2)

# ================================================================ 3 · how it works
s = prs.slides.add_slide(BLANK)
header(s, "How the system works", "Architecture")
textbox(s, 0.6, 1.4, 12.1, 0.45, [
    ("A plain-language prompt is routed through three specialised agents that "
     "mirror a real software team. The Reviewer sends work back until it passes "
     "or a safety limit stops the loop.", 15, False, BLACK)], line_spacing=1.15)

NODES = [
    ("User prompt", "\"Build a snake game\"", LBLUE),
    ("Product Manager", "Spec + numbered acceptance criteria", WHITE),
    ("Developer", "Complete multi-file source", WHITE),
    ("Code Reviewer", "Runs the code, scores the criteria", WHITE),
    ("Finalize", "Code, spec, review history, metrics", LBLUE),
]
bx, bw, bh, gap = 0.72, 2.16, 1.25, 0.27
by = 2.35
for i, (title, sub, fill) in enumerate(NODES):
    x = bx + i * (bw + gap)
    shp = s.shapes.add_shape(5, Inches(x), Inches(by), Inches(bw), Inches(bh))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = NAVY
    shp.line.width = Pt(1.75)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size, r.font.bold, r.font.name = Pt(13.5), True, FONT
    r.font.color.rgb = NAVY
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = sub
    r2.font.size, r2.font.name = Pt(10), FONT
    r2.font.color.rgb = BLACK
    if i < len(NODES) - 1:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.045),
                               Inches(by + bh / 2 - 0.085), Inches(0.18), Inches(0.17))
        a.fill.solid()
        a.fill.fore_color.rgb = NAVY
        a.line.fill.background()
        a.shadow.inherit = False

# feedback arrow beneath Reviewer -> Developer
fy = by + bh + 0.3
dev_x = bx + 2 * (bw + gap) + bw / 2
rev_x = bx + 3 * (bw + gap) + bw / 2
arrow = s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(dev_x), Inches(fy - 0.09),
                           Inches(rev_x - dev_x), Inches(0.2))
arrow.fill.solid()
arrow.fill.fore_color.rgb = RED
arrow.line.fill.background()
arrow.shadow.inherit = False
arrow.adjustments[0] = 0.42
arrow.adjustments[1] = 0.09
for px in (dev_x, rev_x):
    tick = s.shapes.add_shape(1, Inches(px - 0.018), Inches(by + bh), Inches(0.036),
                              Inches(fy - by - bh))
    tick.fill.solid()
    tick.fill.fore_color.rgb = RED
    tick.line.fill.background()
    tick.shadow.inherit = False
textbox(s, dev_x - 0.45, fy + 0.2, 3.6, 0.3,
        [("reject  ∧  iteration < max  ∧  code changed", 11, True, RED)])

card(s, 0.6, 4.55, 3.85, 2.2, "Typed contracts, not prose",
     "Each agent emits structured JSON: the spec, the files, the review verdict. "
     "The next agent reads fields, never English. Malformed output is repaired "
     "once, then recorded as an error rather than crashing the run.", NAVY)
card(s, 4.72, 4.55, 3.85, 2.2, "Everything is recorded",
     "Eleven database tables hold every agent message, every version of the "
     "code, every review and every token count. Any past run can be replayed "
     "and costed.", NAVY)
card(s, 8.85, 4.55, 3.85, 2.2, "Code is executed, not trusted",
     "Generated programs run in a throwaway directory as a separate process "
     "with a wall-clock timeout. A crash or a hang is a result the Reviewer "
     "reads, never an exception that kills the run.", NAVY)
footer(s, 3)

# ============================================== 4 · what was delivered weekly
s = prs.slides.add_slide(BLANK)
header(s, "What was delivered, week by week", "Commit history")
textbox(s, 0.6, 1.4, 12.1, 0.4, [
    ("Each week is a single commit whose message carries that week's milestone "
     "result. The history is the progress record.", 15, False, BLACK)])

table(s, 0.6, 2.05, 12.1, [
    ["Week", "Deliverable", "What it means in practice", "Tests"],
    ["1", "Repository, CI, database schema",
     "Package layout, GitHub Actions on three Python versions, all 11 tables "
     "with indexes and constraints, architecture document", "4"],
    ["2", "State graph and provider layer",
     "LangGraph state machine, four loop-prevention mechanisms, mock and real "
     "LLM adapters behind one interface, runs persisted", "34"],
    ["3", "Product Manager agent",
     "Specification contract with numbered acceptance criteria, validator, "
     "one automatic repair round, five sample prompts", "67"],
    ["4", "Developer agent and sandbox",
     "Complete multi-file output, placeholder rejection, subprocess sandbox "
     "with timeout and path containment, command-line interface", "114"],
], [0.6, 2.9, 7.2, 0.8], size=12.5, row_h=0.84)

textbox(s, 0.6, 6.15, 12.1, 0.8, [
    ("Next:  Week 5 completes the Code Reviewer so it scores the acceptance "
     "criteria by running tests, rather than reporting what the model claims.",
     14, True, NAVY)], line_spacing=1.15)
footer(s, 4)

# ================================================= 5 · measured results
s = prs.slides.add_slide(BLANK)
header(s, "Measured results", "Not estimates")
textbox(s, 0.6, 1.4, 12.1, 0.45, [
    ("Three checks run automatically on every push. Each one executes real "
     "code and compares real output — none of these numbers is an assertion.",
     15, False, BLACK)], line_spacing=1.15)

textbox(s, 0.6, 2.1, 12.1, 0.3,
        [("1 · The loop always terminates, and by the right exit", 15, True, NAVY)])
table(s, 0.6, 2.5, 12.1, [
    ["Scenario", "Iterations", "How it stopped", "Mechanism that fired"],
    ["Reviewer rejects once, then approves", "2", "approved",
     "severity threshold"],
    ["Developer returns identical code", "2", "no_progress",
     "SHA-256 hash of the generated files"],
    ["Reviewer never approves", "5", "iteration_cap", "hard cap, default 5"],
], [4.6, 1.3, 2.2, 4.0], size=12.5, row_h=0.46)

textbox(s, 0.6, 4.45, 5.9, 0.3,
        [("2 · The Product Manager produces usable specs", 15, True, NAVY)])
table(s, 0.6, 4.85, 5.9, [
    ["Check", "Result"],
    ["Valid spec on first answer", "5 / 5"],
    ["Recovered after one repair round", "5 / 5"],
], [4.0, 1.9], size=12.5, row_h=0.44)

textbox(s, 6.8, 4.45, 5.9, 0.3,
        [("3 · The generated code actually runs", 15, True, NAVY)])
table(s, 6.8, 4.85, 5.9, [
    ["Check", "Result"],
    ["Sandbox executions matching expected output", "10 / 10"],
    ["Multi-file project with cross-module imports", "verified"],
], [4.0, 1.9], size=12.5, row_h=0.44)

textbox(s, 0.6, 6.5, 12.1, 0.4, [
    ("114 automated tests, no API key required, zero cost per run.",
     14, True, GREEN)])
footer(s, 5)

# ====================================== 6 · the CS contribution
s = prs.slides.add_slide(BLANK)
header(s, "The core problem: making the loop stop", "Research contribution")
textbox(s, 0.6, 1.4, 12.1, 0.7, [
    ("An agent that reviews another agent's work can reject forever. That is the "
     "failure mode that makes autonomous pipelines unusable, and it is the part "
     "of this project that is genuinely a computer-science problem rather than "
     "an engineering one. Four independent mechanisms guard it.",
     15, False, BLACK)], line_spacing=1.15)

MECH = [
    ("1 · Hard iteration cap",
     "A configurable ceiling, default five, on review-and-fix cycles. The "
     "backstop that always holds."),
    ("2 · Severity threshold",
     "Only critical and major issues block approval. A reviewer that objects to "
     "formatting cannot stall the pipeline."),
    ("3 · No-progress detection",
     "The generated files are hashed. If identical code comes back, the "
     "Developer is not responding to feedback and the loop exits immediately."),
    ("4 · Fallback routing",
     "Exiting without approval does not fail the run. The best version produced "
     "is returned with a warnings report."),
]
for i, (t, b) in enumerate(MECH):
    col, row = i % 2, i // 2
    card(s, 0.6 + col * 6.25, 2.35 + row * 1.6, 5.9, 1.42, t, b, NAVY,
         title_size=14, body_size=12)

band = s.shapes.add_shape(1, Inches(0.6), Inches(5.65), Inches(12.1), Inches(1.05))
band.fill.solid()
band.fill.fore_color.rgb = LBLUE
band.line.color.rgb = NAVY
band.line.width = Pt(1.25)
band.shadow.inherit = False
textbox(s, 0.85, 5.82, 11.6, 0.8, [
    ("Proven, not assumed", 13, True, NAVY),
    ("The routing policy is one pure function of the run state, so an exhaustive "
     "test sweeps every combination of iteration count, severity, verdict and "
     "repeated-or-changed code, and asserts that no reachable state can loop forever.",
     12.5, False, BLACK),
], space_after=3, line_spacing=1.1)
footer(s, 6)

# ============================================ 7 · honest status
s = prs.slides.add_slide(BLANK)
header(s, "What is working — and what is not yet", "Honest status")
textbox(s, 0.6, 1.4, 12.1, 0.35, [
    ("Stated plainly so that the remaining risk is visible rather than discovered later.",
     14, False, BLACK)])

table(s, 0.6, 1.88, 12.1, [
    ["Component", "State", "Evidence or what remains"],
    ["Orchestration — LangGraph state graph", "Working",
     "Runs end to end; 34 integration tests"],
    ["Loop prevention — four mechanisms", "Working",
     "Exhaustive routing sweep; all three exits demonstrated"],
    ["Persistence — 11 tables, SQLite", "Working",
     "Every run queryable; PostgreSQL-ready schema"],
    ["Product Manager agent", "Working",
     "Valid specs on 5/5 sample prompts, both languages"],
    ["Developer agent", "Working",
     "Complete multi-file output; placeholder content rejected"],
    ["Execution sandbox", "Working",
     "10/10 executions; timeout, crash and path escape all handled"],
    ["Code Reviewer agent", "Partial",
     "Parses the contract and calls the sandbox; does not yet score the "
     "criteria by running tests — Week 5"],
    ["Real-model adapters (Anthropic, OpenAI)", "Untested",
     "Written behind the same interface; no live API call made yet — "
     "needs a key, which is a configuration change, not a code change"],
    ["Memory management", "Not started", "Weeks 6–8, as scheduled"],
    ["Benchmark suite and baselines", "Not started", "Weeks 9–12, as scheduled"],
], [4.3, 1.5, 6.3], size=11.5, row_h=0.43, status_col=1,
    status_colors={"Working": GREEN, "Partial": AMBER,
                   "Untested": AMBER, "Not started": GREY})
footer(s, 7)

# ============================================ 8 · how to verify
s = prs.slides.add_slide(BLANK)
header(s, "How to verify any of this", "Four commands")
textbox(s, 0.6, 1.4, 12.1, 0.4, [
    ("Clone the repository and run these. No API key, no network, no cost.",
     15, False, BLACK)])

STEPS = [
    ("pytest -q",
     "Runs the full suite: 114 tests in about eight seconds."),
    ("python -m mawe.cli \"Build a terminal snake game\" --verbose",
     "Watch the pipeline live: spec written, code rejected on a named criterion, "
     "fixed, approved."),
    ("python scripts/m1_demo.py --db",
     "Three scenarios exiting three different ways, and the database rows they "
     "produced."),
    ("python scripts/week4_dev_check.py",
     "Writes generated programs into the sandbox, executes them, compares output."),
]
y = 2.05
for cmd, desc in STEPS:
    chip = s.shapes.add_shape(1, Inches(0.6), Inches(y), Inches(6.6), Inches(0.44))
    chip.fill.solid()
    chip.fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
    chip.line.fill.background()
    chip.shadow.inherit = False
    tfc = chip.text_frame
    tfc.margin_left = Inches(0.14)
    pc = tfc.paragraphs[0]
    pc.alignment = PP_ALIGN.LEFT
    rc = pc.add_run()
    rc.text = cmd
    rc.font.size, rc.font.name = Pt(12), "Consolas"
    rc.font.color.rgb = RGBColor(0x7F, 0xE0, 0x8A)
    textbox(s, 7.4, y + 0.02, 5.3, 0.5, [(desc, 12, False, BLACK)],
            line_spacing=1.08)
    y += 0.78

textbox(s, 0.6, 5.3, 12.1, 0.35,
        [("First-time setup", 13, True, NAVY)])
chip = s.shapes.add_shape(1, Inches(0.6), Inches(5.65), Inches(12.1), Inches(0.44))
chip.fill.solid()
chip.fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
chip.line.fill.background()
chip.shadow.inherit = False
tfc = chip.text_frame
tfc.margin_left = Inches(0.14)
pc = tfc.paragraphs[0]
rc = pc.add_run()
rc.text = ("git clone https://github.com/marcoaperozo/Capstone-MAWE.git  &&  "
           "cd Capstone-MAWE  &&  pip install -r requirements.txt  &&  pip install -e .")
rc.font.size, rc.font.name = Pt(11), "Consolas"
rc.font.color.rgb = RGBColor(0x7F, 0xE0, 0x8A)

textbox(s, 0.6, 6.35, 12.1, 0.6, [
    ("Repository:  github.com/marcoaperozo/Capstone-MAWE", 14, True, NAVY),
    ("Full walkthrough in docs/DEMO-GUIDE.md  ·  design rationale in docs/architecture.md",
     12, False, GREY),
], space_after=2)
footer(s, 8)

prs.save(OUT)
print("saved", OUT)
print(f"{len(prs.slides.__iter__.__self__._sldIdLst)} slides")
